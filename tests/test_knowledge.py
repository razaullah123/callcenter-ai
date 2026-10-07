"""Knowledge base (12.9): text extraction, chunks, the console API, and an agent searching it during a call."""

import base64
import hashlib
import io
import math
import time
import zipfile

import pytest

from runtime.platform import knowledge as kb
from runtime.platform.loader import AgentLoader

from .test_projects import console  # noqa: F401 — the fixture


def fake_vector(text: str) -> list[float]:
    """A deterministic bag-of-words vector (no network): texts sharing words are close."""
    v = [0.0] * kb.EMBED_DIM
    for w in text.lower().split():
        v[int(hashlib.md5(w.strip(".,?!؟").encode()).hexdigest(), 16) % kb.EMBED_DIM] += 1.0
    n = math.sqrt(sum(x * x for x in v)) or 1.0
    return [x / n for x in v]


async def fake_embed(texts: list[str]) -> list[list[float]]:
    return [fake_vector(t) for t in texts]


class FakeProvider:
    async def embed(self, texts):
        return await fake_embed(texts)


@pytest.fixture
def kbconsole(console, monkeypatch):  # noqa: F811
    c, rt, store, loop = console
    from runtime.control import knowledge_api
    c.app.include_router(knowledge_api.router)

    async def embedder(ws):
        return fake_embed
    monkeypatch.setattr(rt.loader, "embedder", embedder)
    monkeypatch.setattr(AgentLoader, "_embedding", staticmethod(lambda providers: FakeProvider()))
    return c, rt, store, loop


def wait_done(c, item_id: str) -> dict:
    for _ in range(100):
        item = c.get(f"/api/knowledge/{item_id}").json()
        if item["status"] != "processing":
            return item
        time.sleep(0.05)
    raise AssertionError("still processing")


def test_text_extraction():
    assert kb.extract_text("a.txt", "Visiting hours: 9 to 5".encode())[0] == "Visiting hours: 9 to 5"
    html = b"<html><head><style>x{}</style></head><body><h1>Parking</h1><p>Free &amp; open 24h</p></body></html>"
    text, ext = kb.extract_text("page.HTM", html)
    assert ext == "html" and "Parking" in text and "Free & open 24h" in text and "x{}" not in text
    import docx
    d = docx.Document()
    d.add_paragraph("مواعيد الزيارة من ٩ إلى ٥")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Clinic", "Dental"
    buf = io.BytesIO()
    d.save(buf)
    text, ext = kb.extract_text("faq.docx", buf.getvalue())
    assert ext == "docx" and "مواعيد الزيارة" in text and "Clinic | Dental" in text
    z = io.BytesIO()
    with zipfile.ZipFile(z, "w") as f:
        f.writestr("OEBPS/ch1.xhtml", "<html><body><p>Chapter one text</p></body></html>")
    assert "Chapter one text" in kb.extract_text("book.epub", z.getvalue())[0]
    for name, msg in (("old.doc", ".docx"), ("x.exe", "unsupported"), ("empty.txt", "no text")):
        with pytest.raises(ValueError, match=msg):
            kb.extract_text(name, b"" if name == "empty.txt" else b"x")


def test_chunks_are_bounded_and_overlap():
    para = " ".join(f"Sentence number {i} about the clinic." for i in range(60))
    parts = kb.chunk(para + "\n\nShort tail paragraph.")
    assert len(parts) > 2 and all(len(p) <= kb.CHUNK_CHARS + kb.CHUNK_OVERLAP + 40 for p in parts)
    assert parts[0][-20:].split()[-1] in parts[1]                      # overlap carries context over
    assert kb.chunk("one line") == ["one line"]


def test_knowledge_api_and_search(kbconsole):
    c, rt, store, loop = kbconsole
    r = c.post("/api/knowledge/text", json={"name": "Parking", "content": "Parking is free for patients in the basement."})
    assert r.status_code == 200, r.text
    text_id = r.json()["id"]
    doc = "Visiting hours are from 9 am to 9 pm.\n\nChildren under 12 can't visit the ICU."
    r = c.post("/api/knowledge/file", json={"filename": "visits.txt", "data": base64.b64encode(doc.encode()).decode()})
    assert r.status_code == 200, r.text
    file_id = r.json()["id"]
    assert wait_done(c, text_id)["status"] == "completed" and wait_done(c, file_id)["chunks"] == 1
    listing = c.get("/api/knowledge").json()
    assert {i["id"] for i in listing["items"]} == {text_id, file_id} and listing["usage_bytes"] > 0
    assert "content" not in listing["items"][0]                       # the list doesn't carry the texts
    hits = c.post("/api/knowledge/search", json={"query": "when are visiting hours"}).json()
    assert hits["mode"] == "hybrid" and "Visiting hours" in hits["results"][0]["text"]
    assert c.post("/api/knowledge/file", json={"filename": "x.doc", "data": "eA=="}).status_code == 422
    # a scope check: another project can't see it
    assert c.get(f"/api/knowledge/{text_id}", headers={"X-Project": "nope"}).status_code in (403, 404)
    assert c.delete(f"/api/knowledge/{text_id}").json()["deleted"] == text_id
    assert c.get(f"/api/knowledge/{text_id}").status_code == 404


def test_keyword_search_without_embeddings(kbconsole):
    c, rt, store, loop = kbconsole

    async def none(ws):
        return None
    rt.loader.embedder = none
    item = c.post("/api/knowledge/text", json={"name": "Hours", "content": "The pharmacy opens at 8 am."}).json()
    done = wait_done(c, item["id"])
    assert done["status"] == "completed_with_errors" and "keyword" in done["error"]
    res = loop.run_until_complete(kb.search(store, "hmg", [item["id"]], "pharmacy", None))
    assert res["mode"] == "keyword" and "pharmacy" in res["results"][0]["text"]


def test_an_agent_searches_its_knowledge(kbconsole):
    c, rt, store, loop = kbconsole
    item = c.post("/api/knowledge/text", json={"name": "ICU", "content": "ICU visiting is limited to two visitors."}).json()
    wait_done(c, item["id"])
    r = c.post("/api/knowledge/use", json={"agent": "hmg-care", "items": [item["id"]]})
    assert r.status_code == 200, r.text
    row = c.get("/api/knowledge").json()["items"][0]
    assert row["draft_by"][0]["agent"] == "hmg-care" and not row["used_by"]            # live after Publish
    agent = loop.run_until_complete(rt.loader.load_draft("hmg-care"))
    assert "search_knowledge_base" in agent.executor.catalog.tools
    assert "search_knowledge_base" in agent.skills.global_tools
    from runtime.events import EventBus
    from runtime.providers import ToolCall
    from runtime.tools import ToolContext

    async def run():
        bus = EventBus()
        await bus.start()
        try:
            return await agent.executor.execute(ToolCall(id="t1", name="search_knowledge_base",
                                                         arguments={"query": "how many ICU visitors"}),
                                                ToolContext(call_id="c1"), bus.bind(call_id="c1"))
        finally:
            await bus.stop()
    result = loop.run_until_complete(run())
    assert result.ok and "two visitors" in str(result.data)
    # removing every item takes the tool away again
    assert c.post("/api/knowledge/use", json={"agent": "hmg-care", "items": []}).json()["items"] == []
    agent = loop.run_until_complete(rt.loader.load_draft("hmg-care"))
    assert "search_knowledge_base" not in agent.executor.catalog.tools
    assert c.post("/api/knowledge/use", json={"agent": "hmg-care", "items": ["kb_nope"]}).status_code == 404


# ---------------------------------------------------------------- Hamsa parity: URL items, rename, delete protection


@pytest.fixture
def web(monkeypatch):
    """A fake internet: {url: (status, content type, body)}; every host counts as public except *.internal."""
    import httpx
    pages: dict[str, tuple] = {}
    seen: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(str(request.url))
        status, kind, body = pages.get(str(request.url), (404, "text/plain", "not found"))
        return httpx.Response(status, headers={"content-type": kind, **({"location": body} if status in (301, 302) else {})},
                              content=b"" if status in (301, 302) else body.encode())

    async def public(host):
        return not host.endswith(".internal")
    monkeypatch.setattr(kb, "_client", httpx.AsyncClient(transport=httpx.MockTransport(handler)))
    monkeypatch.setattr(kb, "_public_host", public)
    pages["seen"] = seen          # type: ignore[assignment]
    return pages


def test_url_checks(web):
    import asyncio
    run = asyncio.new_event_loop().run_until_complete
    assert run(kb.check_url("https://example.com/faq")) == "https://example.com/faq"
    for bad, msg in (("http://example.com", "https"), ("https://localhost/x", "https"), ("https://example.com/a.mp4", "media"),
                     ("https://user:pw@example.com/", "user name"), ("https://db.internal/x", "not a public")):
        with pytest.raises(ValueError, match=msg):
            run(kb.check_url(bad))


def test_adding_a_web_page(kbconsole, web):
    c, rt, store, loop = kbconsole
    web["https://example.com/faq"] = (200, "text/html", "<html><body><h1>Visiting</h1><p>Open 9 to 9.</p><script>x()</script></body></html>")
    r = c.post("/api/knowledge/url", json={"name": "FAQ page", "url": "https://example.com/faq"})
    assert r.status_code == 200, r.text
    item = wait_done(c, r.json()["id"])
    assert item["status"] == "completed" and item["type"] == "url" and item["source_url"] == "https://example.com/faq"
    assert "Open 9 to 9." in item["content"] and "x()" not in item["content"]
    hits = c.post("/api/knowledge/search", json={"query": "open 9"}).json()
    assert hits["results"] and "Open 9" in hits["results"][0]["text"]


def test_a_page_that_cannot_be_read_fails_with_the_reason(kbconsole, web):
    c, rt, store, loop = kbconsole
    web["https://example.com/private"] = (403, "text/html", "no")
    item = wait_done(c, c.post("/api/knowledge/url", json={"name": "P", "url": "https://example.com/private"}).json()["id"])
    assert item["status"] == "failed" and "refused access" in item["error"]
    # refused up front: not https, a private host
    assert c.post("/api/knowledge/url", json={"name": "P", "url": "http://example.com/"}).status_code == 422
    assert c.post("/api/knowledge/url", json={"name": "P", "url": "https://db.internal/"}).status_code == 422


def test_redirect_to_a_private_host_is_refused(kbconsole, web):
    c, rt, store, loop = kbconsole
    web["https://example.com/go"] = (302, "text/html", "https://db.internal/secret")
    web["https://db.internal/secret"] = (200, "text/html", "<p>secret</p>")
    item = wait_done(c, c.post("/api/knowledge/url", json={"name": "R", "url": "https://example.com/go"}).json()["id"])
    assert item["status"] == "failed" and "not a public" in item["error"]
    assert "https://db.internal/secret" not in web["seen"]


def test_sitemap_discovery_and_a_multi_page_item(kbconsole, web):
    c, rt, store, loop = kbconsole
    locs = "".join(f"<url><loc>https://example.com/p{i}</loc></url>" for i in range(130))
    web["https://example.com/sitemap.xml"] = (200, "application/xml", f"<urlset>{locs}<url><loc>https://other.com/x</loc></url>"
                                                                       "<url><loc>https://example.com/logo.png</loc></url></urlset>")
    urls = c.post("/api/knowledge/url/discover", json={"url": "https://example.com/"}).json()["urls"]
    assert len(urls) == kb.MAX_URLS and urls[0] == "https://example.com/p0" and all(u.startswith("https://example.com/") for u in urls)
    assert c.post("/api/knowledge/url/discover", json={"url": "https://nowhere.example/"}).json() == {"urls": []}
    web["https://example.com/p0"] = (200, "text/html", "<p>Zero page</p>")
    web["https://example.com/p1"] = (200, "text/html", "<p>First page</p>")
    r = c.post("/api/knowledge/url", json={"name": "Site", "url": "https://example.com/", "urls": [
        "https://example.com/p0", "https://example.com/p1", "https://example.com/p2"]})        # p2 is a 404
    item = wait_done(c, r.json()["id"])
    assert item["status"] == "completed_with_errors" and "1 of 3 pages could not be read" in item["error"]
    assert "Source: https://example.com/p0\nZero page" in item["content"] and "First page" in item["content"]
    assert c.post("/api/knowledge/url", json={"name": "Big", "url": "https://example.com/",
                                              "urls": [f"https://example.com/p{i}" for i in range(101)]}).status_code == 422


def test_rename_keeps_the_content(kbconsole):
    c, rt, store, loop = kbconsole
    item = c.post("/api/knowledge/text", json={"name": "Old", "content": "Parking is free."}).json()
    wait_done(c, item["id"])
    assert c.patch(f"/api/knowledge/{item['id']}", json={"name": "  Parking policy "}).json()["name"] == "Parking policy"
    got = c.get(f"/api/knowledge/{item['id']}").json()
    assert got["name"] == "Parking policy" and got["content"] == "Parking is free." and got["status"] == "completed"
    assert c.patch(f"/api/knowledge/{item['id']}", json={"name": "   "}).status_code == 422
    assert c.patch("/api/knowledge/kb_missing", json={"name": "x"}).status_code == 404


def test_an_item_in_use_cannot_be_deleted(kbconsole):
    c, rt, store, loop = kbconsole
    item = c.post("/api/knowledge/text", json={"name": "ICU", "content": "ICU visiting is limited."}).json()
    wait_done(c, item["id"])
    assert c.post("/api/knowledge/use", json={"agent": "hmg-care", "items": [item["id"]]}).status_code == 200
    r = c.delete(f"/api/knowledge/{item['id']}")
    assert r.status_code == 409 and "ICU is used by" in r.json()["detail"]
    c.post("/api/knowledge/use", json={"agent": "hmg-care", "items": []})
    assert c.delete(f"/api/knowledge/{item['id']}").status_code == 200
