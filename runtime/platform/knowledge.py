"""Knowledge base (12.9): documents and free text an agent searches during a call.

Adding an item: its text is extracted (TXT / MD / HTML / PDF / DOCX / EPUB), split into ~800-character chunks on
paragraph and sentence boundaries, and each chunk embedded with the project's embedding connection (pgvector,
1024 dims). Chunks whose embedding fails keep a keyword index, so search still works — the item is then
"completed_with_errors".

Searching (the `search_knowledge_base` tool an agent gets when its release lists knowledge items): the question is
embedded and matched by vector similarity and by keywords, fused by reciprocal rank (`store.kb_search`). When the
embedding service is down the search is keyword-only.
"""

import asyncio
import html
import io
import logging
import re
import time
import zipfile
from html.parser import HTMLParser
from typing import Any, Awaitable, Callable

log = logging.getLogger(__name__)

MAX_FILE_BYTES = 21 * 1024 * 1024          # per document (as Hamsa)
QUOTA_BYTES = 300 * 1024 * 1024            # per project
MAX_TEXT_CHARS = 25_000                    # free text items
EXTENSIONS = ("pdf", "docx", "txt", "md", "html", "htm", "epub")
CHUNK_CHARS, CHUNK_OVERLAP = 800, 120
EMBED_DIM = 1024
EMBED_BATCH = 16
SEARCH_EMBED_TIMEOUT_S = 3.0

Embedder = Callable[[list[str]], Awaitable[list[list[float]]]]


# ---------------------------------------------------------------- text extraction

class _Text(HTMLParser):
    SKIP = {"script", "style", "head", "noscript"}
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "table"}

    def __init__(self) -> None:
        super().__init__()
        self.out: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skip:
            self.skip -= 1
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.out.append(data)


def html_text(markup: str) -> str:
    p = _Text()
    p.feed(markup)
    return html.unescape("".join(p.out))


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1256", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def extract_text(filename: str, data: bytes) -> tuple[str, str]:
    """(text, extension) of a document; ValueError when the type isn't supported or nothing can be read."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "doc":
        raise ValueError("old Word .doc files aren't supported — save it as .docx or PDF")
    if ext not in EXTENSIONS:
        raise ValueError(f"unsupported file type .{ext or '?'} — use PDF, DOCX, TXT, MD, HTML or EPUB")
    if ext in ("txt", "md"):
        text = _decode(data)
    elif ext in ("html", "htm"):
        text = html_text(_decode(data))
    elif ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("the PDF is password-protected") from None
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    elif ext == "docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
    else:                                                  # epub: a zip of XHTML chapters
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            names = sorted(n for n in z.namelist() if n.lower().endswith((".xhtml", ".html", ".htm")))
            text = "\n\n".join(html_text(_decode(z.read(n))) for n in names)
        ext = "epub"
    text = normalize(text)
    if not text:
        raise ValueError("no text found in the file (a scanned PDF needs OCR first)")
    return text, "html" if ext == "htm" else ext


def normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def words(text: str) -> int:
    return len(text.split())


# ---------------------------------------------------------------- web pages (url items)

MAX_URLS = 100                              # pages per url item (a sitemap), as Hamsa
MAX_PAGE_BYTES = 5 * 1024 * 1024
FETCH_TIMEOUT_S = 15.0
MEDIA_EXT = (".mp3", ".mp4", ".wav", ".avi", ".mov", ".webm", ".jpg", ".jpeg", ".png", ".gif", ".svg", ".webp", ".zip")
_client: Any = None


def _http():
    global _client
    if _client is None:
        import httpx
        _client = httpx.AsyncClient(timeout=FETCH_TIMEOUT_S, follow_redirects=False,
                                    headers={"User-Agent": "Mozilla/5.0 (compatible; VoiceAgentKB/1.0)"})
    return _client


async def _public_host(host: str) -> bool:
    """True when `host` resolves only to public addresses (nothing on our own network is fetched)."""
    import ipaddress
    import socket
    try:
        infos = await asyncio.get_running_loop().getaddrinfo(host, 443, type=socket.SOCK_STREAM)
    except OSError:
        return False
    for info in infos:
        ip = ipaddress.ip_address(info[4][0].split("%")[0])
        if not ip.is_global:
            return False
    return bool(infos)


async def check_url(url: str) -> str:
    """The URL normalised, or ValueError when we must not fetch it (not https, a media file, a private address)."""
    from urllib.parse import urlsplit
    u = urlsplit((url or "").strip())
    if u.scheme != "https" or not u.hostname or "." not in u.hostname:
        raise ValueError("only public https:// links are supported")
    if u.username or u.password:
        raise ValueError("links with a user name or password are not supported")
    if len(url) > 2048:
        raise ValueError("the link is too long")
    if u.path.lower().endswith(MEDIA_EXT):
        raise ValueError("media files (audio, video, images) are not supported")
    if not await _public_host(u.hostname):
        raise ValueError(f"{u.hostname} is not a public website")
    return u.geturl()


async def _get(url: str, hops: int = 3) -> tuple[str, str]:
    """(content type, text) of a public page; follows up to `hops` redirects, checking every hop."""
    from urllib.parse import urljoin
    for _ in range(hops + 1):
        url = await check_url(url)
        async with _http().stream("GET", url) as r:
            if r.is_redirect and r.headers.get("location"):
                url = urljoin(url, r.headers["location"])
                continue
            if r.status_code in (401, 403):
                raise ValueError(f"the site refused access (HTTP {r.status_code}) — it needs a login or blocks scraping")
            if not r.is_success:
                raise ValueError(f"the page answered HTTP {r.status_code}")
            kind = r.headers.get("content-type", "").split(";")[0].strip().lower()
            body = b""
            async for part in r.aiter_bytes():
                body += part
                if len(body) > MAX_PAGE_BYTES:
                    raise ValueError(f"the page is larger than {MAX_PAGE_BYTES // (1024 * 1024)} MB")
            return kind, _decode(body)
    raise ValueError("too many redirects")


async def fetch_page_text(url: str) -> str:
    kind, body = await _get(url)
    if kind and not (kind.startswith("text/") or "html" in kind or "xml" in kind):
        raise ValueError(f"unsupported content type {kind}")
    text = normalize(body if kind == "text/plain" else html_text(body))
    if not text:
        raise ValueError("no text found on the page (it may be built by JavaScript)")
    return text


async def discover_urls(url: str) -> list[str]:
    """Pages the site's sitemap lists (same host, https, at most MAX_URLS); [] when it has none."""
    from urllib.parse import urlsplit
    start = urlsplit(await check_url(url))
    origin = f"https://{start.netloc}"
    found: list[str] = []
    queue = [f"{origin}/sitemap.xml"]
    seen_maps: set[str] = set()
    while queue and len(found) < MAX_URLS and len(seen_maps) < 5:
        sm = queue.pop(0)
        if sm in seen_maps:
            continue
        seen_maps.add(sm)
        try:
            _, xml = await _get(sm)
        except Exception:                                    # noqa: BLE001 — no sitemap here
            continue
        for loc in re.findall(r"<loc>\s*([^<\s]+)\s*</loc>", xml):
            loc = html.unescape(loc)
            if urlsplit(loc).netloc != start.netloc or not loc.startswith("https://"):
                continue
            if loc.lower().endswith(".xml"):
                queue.append(loc)
            elif not loc.lower().endswith(MEDIA_EXT) and loc not in found:
                found.append(loc)
                if len(found) >= MAX_URLS:
                    break
    return found


async def fetch_site_text(urls: list[str]) -> tuple[str, list[str]]:
    """(text of the pages that could be read, one error line per page that couldn't). ValueError if none could."""
    sem = asyncio.Semaphore(4)

    async def one(u: str):
        async with sem:
            try:
                return u, await fetch_page_text(u), None
            except Exception as e:                           # noqa: BLE001
                return u, None, str(e) or e.__class__.__name__
    got = await asyncio.gather(*(one(u) for u in urls))
    ok = [(u, t) for u, t, _ in got if t]
    errors = [f"{u}: {e}" for u, _, e in got if e]
    if not ok:
        raise ValueError(errors[0].split(": ", 1)[1] if len(urls) == 1 else "none of the pages could be read — " + errors[0])
    text = "\n\n".join(f"Source: {u}\n{t}" if len(urls) > 1 else t for u, t in ok)
    return text, errors


# ---------------------------------------------------------------- chunks

_SENTENCE = re.compile(r"(?<=[.!?؟。])\s+")


def chunk(text: str, size: int = CHUNK_CHARS, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """~`size`-character chunks on paragraph / sentence boundaries; consecutive chunks share ~`overlap` chars."""
    pieces: list[str] = []
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        if len(para) <= size:
            pieces.append(para)
            continue
        for sent in _SENTENCE.split(para):
            while len(sent) > size:                      # a very long sentence: hard split
                pieces.append(sent[:size])
                sent = sent[size - overlap:]
            if sent.strip():
                pieces.append(sent.strip())
    chunks: list[str] = []
    cur = ""
    for p in pieces:
        if cur and len(cur) + 1 + len(p) > size:
            chunks.append(cur)
            tail = cur[-overlap:]
            cur = (tail[tail.find(" ") + 1:] if " " in tail else tail) + " " + p
        else:
            cur = f"{cur}\n{p}" if cur else p
    if cur.strip():
        chunks.append(cur)
    return chunks


# ---------------------------------------------------------------- ingest / search

async def ingest(store, item_id: str, ws: str, text: str, embed: Embedder | None) -> dict[str, Any]:
    """Chunk + embed `text` into the item's chunks and finish its status. Never raises (status says what failed)."""
    started = time.monotonic()
    try:
        parts = chunk(text)
        vectors: list[list[float] | None] = [None] * len(parts)
        failed = 0
        if embed is not None:
            for i in range(0, len(parts), EMBED_BATCH):
                batch = parts[i:i + EMBED_BATCH]
                try:
                    got = await embed(batch)
                    for j, v in enumerate(got):
                        vectors[i + j] = v if v and len(v) == EMBED_DIM else None
                        failed += vectors[i + j] is None
                except Exception as e:                    # noqa: BLE001 — keyword search still works
                    failed += len(batch)
                    log.warning("knowledge %s: embedding failed for chunks %d-%d: %r", item_id, i, i + len(batch), e)
        else:
            failed = len(parts)
        await store.set_kb_chunks(item_id, ws, [(n, t, vectors[n]) for n, t in enumerate(parts)])
        status = "completed" if not failed else "completed_with_errors"
        error = None if not failed else (f"{failed} of {len(parts)} parts have no semantic index (embedding service "
                                         "unavailable) — keyword search only for them")
        await store.put_kb_item({"id": item_id, "status": status, "error": error, "chunks": len(parts),
                                 "words": words(text)})
        return {"status": status, "chunks": len(parts), "ms": round((time.monotonic() - started) * 1000)}
    except Exception as e:                               # noqa: BLE001
        log.exception("knowledge %s: ingest failed", item_id)
        await store.put_kb_item({"id": item_id, "status": "failed", "error": str(e)[:500]})
        return {"status": "failed", "error": str(e)}


async def search(store, ws: str, item_ids: list[str], query: str, embed: Embedder | None, k: int = 4) -> dict:
    query = (query or "").strip()
    if not query or not item_ids:
        return {"results": [], "note": "nothing to search"}
    vector = None
    mode = "keyword"
    if embed is not None:
        try:
            got = await asyncio.wait_for(embed([query]), SEARCH_EMBED_TIMEOUT_S)
            if got and len(got[0]) == EMBED_DIM:
                vector, mode = got[0], "hybrid"
        except Exception as e:                           # noqa: BLE001
            log.warning("knowledge search: embedding unavailable, keyword only: %r", e)
    rows = await store.kb_search(ws, item_ids, query, vector, k)
    return {"mode": mode, "results": [{"source": r["name"], "text": r["text"]} for r in rows]}


def embedder_from(provider) -> Embedder | None:
    """An embedding provider (runtime.providers) as an Embedder, or None."""
    if provider is None:
        return None

    async def embed(texts: list[str]) -> list[list[float]]:
        return await provider.embed(texts)
    return embed


__all__ = ["MAX_URLS", "check_url", "discover_urls", "fetch_page_text", "fetch_site_text", "EXTENSIONS", "MAX_FILE_BYTES", "MAX_TEXT_CHARS", "QUOTA_BYTES", "chunk", "embedder_from",
           "extract_text", "ingest", "normalize", "search", "words"]
